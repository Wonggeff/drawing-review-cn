# -*- coding: utf-8 -*-
"""Generate the formal drawing-review record workbook (xlsx) and document (docx)."""
import os, re, csv, sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT

import os, sys as _sys
_sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import cfg  # 路径配置见 common.py / config.json
ROOT = cfg.out_dir
FIND = cfg.findings_dir
MD02 = cfg.main_issue_md

# ---------------- 解析 02 中的 Q / X / 需确认事项 ----------------
txt = open(MD02, encoding="utf-8").read()

def clean(s):
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = s.replace("`", "").replace("｜", "|")
    return s.strip()

def blocks(pattern):
    out = []
    ms = list(re.finditer(pattern, txt, re.M))
    for i, m in enumerate(ms):
        start = m.end()
        nxt = len(txt)
        for j in range(i + 1, len(ms)):
            nxt = ms[j].start()
            break
        if not ms[i + 1:]:
            nxt = len(txt)
        body = txt[start:nxt]
        body = re.split(r"\n---\s*\n", body)[0]
        out.append((m, body))
    return out

Q = []
for m, body in blocks(r"^### (Q-\d+)【([^·】]+)·([^】]+)】(.+?)$"):
    body_txt = clean(body)
    # 提取涉及图纸行
    dwg = ""
    mm = re.search(r"\*\*涉及图纸\*\*：(.+)", body)
    if mm:
        dwg = clean(mm.group(1))
    else:
        mm = re.search(r"涉及图纸：(.+)", clean(body))
        dwg = mm.group(1) if mm else ""
    Q.append({
        "id": m.group(1), "cls": m.group(2).strip(), "sev": m.group(3).strip(),
        "title": clean(m.group(4)), "dwg": dwg, "body": body_txt,
    })

X = []
for m, body in blocks(r"^### (X-\d+) (.+?)$"):
    X.append({"id": m.group(1), "title": clean(m.group(2)), "body": clean(body)})

# 需设计确认事项表
CONF = []
sec = txt.split("## 6. 需设计单位书面确认／补充事项汇总")
if len(sec) > 1:
    for line in sec[1].splitlines():
        if not line.strip().startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) < 4 or cells[0] in ("序号",) or set(cells[0]) <= set("-: "):
            continue
        CONF.append(cells[:4])

DET = list(csv.DictReader(open(os.path.join(FIND, "_merged.csv"), encoding="utf-8-sig")))

print("Q=%d X=%d CONF=%d DET=%d" % (len(Q), len(X), len(CONF), len(DET)))

# ---------------- Excel ----------------
THIN = Side(style="thin", color="9E9E9E")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="微软雅黑", size=9)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

# Sheet 0: 说明
ws = wb.active
ws.title = "说明"
rows = [
    ["图纸会审记录表"], [""],
    ["工程名称", cfg.project_name], ["建设地点", cfg.project_location],
    ["建设单位", cfg.client], ["设计单位", cfg.designer],
    ["勘察单位", cfg.surveyor], ["施工图审查机构", cfg.reviewer],
    ["会审时间", ""], ["会审地点", ""], ["主持人", ""],
    ["参加单位及人员", "建设单位：　　　　　设计单位：　　　　　监理单位：　　　　　施工单位：　　　　　其他：　　　　　"],
    [""],
    ["问题总数", len(DET)], ["其中 A 设计错误", sum(1 for d in DET if d["类别"].strip().startswith("A"))],
    ["其中 B 图纸矛盾", sum(1 for d in DET if d["类别"].strip().startswith("B"))],
    ["其中 C 深度不足", sum(1 for d in DET if d["类别"].strip().startswith("C"))],
    ["其中 D 施工建议", sum(1 for d in DET if d["类别"].strip().startswith("D"))],
    ["严重程度 高", sum(1 for d in DET if "高" in d["严重程度"])],
    ["严重程度 中", sum(1 for d in DET if "中" in d["严重程度"])],
    ["严重程度 低", sum(1 for d in DET if "低" in d["严重程度"])],
    [""],
    ["填写说明", "1. “设计回复”“处理结果”“闭环状态”三列为空白待填写列；“闭环状态”已设下拉（未闭合/已闭合/作废/转设计变更）。"],
    ["", "2. 问题级别：A 设计错误（须出设计变更，涉及强制性标准须重新报审）／B 图纸矛盾（须书面确认）／C 深度不足（须补充说明或详图）／D 施工建议（设计确认后执行）。"],
    ["", "3. 凡标注“需设计确认”的条目，须取得设计单位书面答复后方可施工；涉及结构安全与消防的条目应在会审会上形成明确结论。"],
    ["", "4. 本表共 6 个工作表：说明、问题明细（703 条）、一级问题、跨专业矛盾、需设计确认事项、统计。"],
]
for r in rows:
    ws.append(r)
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 110
ws["A1"].font = Font(name="微软雅黑", size=16, bold=True)
for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=2):
    for c in row:
        c.font = CELL_FONT
        c.alignment = WRAP
    row[0].font = Font(name="微软雅黑", size=9, bold=True)

# Sheet 1: 问题明细
ws2 = wb.create_sheet("问题明细(703条)")
HEAD = ["序号", "统一编号", "专业", "审查范围", "图号/图名", "位置（轴线/标高/房间）", "问题描述",
        "类别", "严重程度", "建议处理措施", "依据", "提出单位", "设计回复", "处理结果", "闭环状态"]
ws2.append(HEAD)
for i, d in enumerate(DET, 1):
    ws2.append([i, d["统一编号"], d["专业"], d["审查范围"], d["图号图名"], d["位置"], d["问题描述"],
                d["类别"], d["严重程度"], d["建议处理措施"], d["依据"], "施工单位", "", "", ""])
WID = [6, 11, 12, 22, 34, 26, 70, 8, 9, 40, 26, 10, 26, 20, 12]
for i, w in enumerate(WID, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for c in ws2[1]:
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=len(HEAD)):
    for c in row:
        c.font = CELL_FONT; c.alignment = WRAP; c.border = BORDER
ws2.freeze_panes = "C2"
ws2.auto_filter.ref = "A1:O%d" % ws2.max_row
dv = DataValidation(type="list", formula1='"未闭合,已闭合,作废,转设计变更"', allow_blank=True)
ws2.add_data_validation(dv)
dv.add("O2:O%d" % ws2.max_row)

# Sheet 2: 一级问题
ws3 = wb.create_sheet("一级问题")
ws3.append(["编号", "类别", "严重程度", "问题标题", "涉及图纸", "问题描述与依据", "建议处理措施", "设计回复", "处理结果"])
for q in Q:
    # 拆分"建议处理"段
    body = q["body"]
    m = re.search(r"\*\*建议处理\*\*：(.*)", body, re.S)
    sug = clean(m.group(1)) if m else ""
    desc = body.split("**建议处理**")[0]
    desc = re.sub(r"^\s*", "", desc)
    ws3.append([q["id"], q["cls"], q["sev"], q["title"], q["dwg"], desc, sug, "", ""])
for i, w in enumerate([8, 10, 10, 34, 40, 90, 50, 24, 16], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
for c in ws3[1]:
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, max_col=9):
    for c in row:
        c.font = CELL_FONT; c.alignment = WRAP; c.border = BORDER
ws3.freeze_panes = "C2"

# Sheet 3: 跨专业矛盾
ws4 = wb.create_sheet("跨专业矛盾")
ws4.append(["编号", "议题", "问题描述与依据", "建议处理措施", "设计回复", "处理结果"])
for x in X:
    body = x["body"]
    m = re.search(r"\*\*须(?:要求|核实|由)[^\n]*", body)
    ws4.append([x["id"], x["title"], body, "", "", ""])
for i, w in enumerate([8, 30, 110, 40, 24, 16], 1):
    ws4.column_dimensions[get_column_letter(i)].width = w
for c in ws4[1]:
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
for row in ws4.iter_rows(min_row=2, max_row=ws4.max_row, max_col=6):
    for c in row:
        c.font = CELL_FONT; c.alignment = WRAP; c.border = BORDER
ws4.freeze_panes = "B2"

# Sheet 4: 需设计确认事项
ws5 = wb.create_sheet("需设计确认事项")
ws5.append(["序号", "事项", "对应问题编号", "要求交付物", "设计回复", "交付日期", "完成情况"])
for r in CONF:
    ws5.append([r[0], r[1], r[2], r[3], "", "", ""])
for i, w in enumerate([6, 44, 14, 44, 24, 12, 12], 1):
    ws5.column_dimensions[get_column_letter(i)].width = w
for c in ws5[1]:
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
for row in ws5.iter_rows(min_row=2, max_row=ws5.max_row, max_col=7):
    for c in row:
        c.font = CELL_FONT; c.alignment = WRAP; c.border = BORDER
ws5.freeze_panes = "B2"
dv2 = DataValidation(type="list", formula1='"未提交,部分提交,已提交,已确认"', allow_blank=True)
ws5.add_data_validation(dv2)
dv2.add("G2:G%d" % ws5.max_row)

# Sheet 5: 统计
ws6 = wb.create_sheet("统计")
ws6.append(["审查范围（分册）", "条数", "A设计错误", "B图纸矛盾", "C深度不足", "D施工建议", "高", "中", "低"])
from collections import OrderedDict
srcs = OrderedDict()
for d in DET:
    k = d["审查范围"]
    s = srcs.setdefault(k, dict(n=0, A=0, B=0, C=0, D=0, 高=0, 中=0, 低=0))
    s["n"] += 1
    c0 = (d["类别"] or "")[:1]
    if c0 in "ABCD":
        s[c0] += 1
    for kk in ("高", "中", "低"):
        if kk in (d["严重程度"] or ""):
            s[kk] += 1
for k, v in srcs.items():
    ws6.append([k, v["n"], v["A"], v["B"], v["C"], v["D"], v["高"], v["中"], v["低"]])
tot = ["合计", len(DET)] + [sum(v[x] for v in srcs.values()) for x in "ABCD高 中低".replace(" ", "")]
ws6.append(tot)
for i, w in enumerate([30, 8, 11, 11, 11, 11, 8, 8, 8], 1):
    ws6.column_dimensions[get_column_letter(i)].width = w
for c in ws6[1]:
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
for row in ws6.iter_rows(min_row=2, max_row=ws6.max_row, max_col=9):
    for c in row:
        c.font = CELL_FONT; c.alignment = CENTER; c.border = BORDER
for c in ws6[ws6.max_row]:
    c.font = Font(name="微软雅黑", size=9, bold=True)

XLSX = os.path.join(ROOT, "04_图纸会审记录表.xlsx")
wb.save(XLSX)
print("saved", XLSX)

# ---------------- Word ----------------
doc = Document()
st = doc.styles["Normal"]
st.font.name = "宋体"
st.font.size = Pt(9)
st._element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "宋体")

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.2)
sec.top_margin = sec.bottom_margin = Cm(1.2)

def H(text, size=14, align="center", space=6):
    p = doc.add_paragraph()
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    r.font.name = "黑体"; r._element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "黑体")
    p.paragraph_format.space_before = Pt(space); p.paragraph_format.space_after = Pt(space)
    return p

def P(text, size=9, align="left"):
    p = doc.add_paragraph()
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    r = p.add_run(text); r.font.size = Pt(size)
    return p

def mktable(headers, widths, rows, fsize=7.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); r.bold = True; r.font.size = Pt(fsize)
        r.font.name = "黑体"; r._element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "黑体")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v)); r.font.size = Pt(fsize)
            r.font.name = "宋体"; r._element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "宋体")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return t

H(cfg.project_name, 18)
H("施工图图纸会审记录", 16)
P("")
info = [
    ("工程名称", cfg.project_name, "会审时间", ""),
    ("建设地点", cfg.project_location, "会审地点", ""),
    ("建设单位", cfg.client, "主持人", ""),
    ("设计单位", cfg.designer, "记录人", ""),
    ("勘察单位", cfg.surveyor, "审查机构", cfg.reviewer),
    ("施工图审查合格章", "", "问题总条数", str(len(DET))),
]
mktable(["项目", "内容", "项目", "内容"], [3.0, 9.5, 3.0, 9.5], info, fsize=9)

P("")
P("参加会审单位及人员（签字）：", 10)
mktable(["单位类别", "单位名称", "项目负责人/专业负责人（签字）", "盖章"],
        [3.5, 8.0, 8.5, 5.0],
        [["建设单位", cfg.client, "", ""],
         ["设计单位", cfg.designer, "", ""],
         ["勘察单位", cfg.surveyor, "", ""],
         ["监理单位", "", "", ""],
         ["施工单位", "", "", ""]], fsize=9)
P("")
P("会审结论：", 10)
P("本次施工图自审共提出问题 %d 条（A 设计错误 %d 条、B 图纸矛盾 %d 条、C 深度不足 %d 条、D 施工建议 %d 条），"
  "归纳一级问题 %d 项、跨专业矛盾 %d 项、需设计单位书面确认／补充事项 %d 项。详细明细见附件《04_图纸会审记录表.xlsx》"
  "及《02_图纸会审问题清单.md》《03_各专业问题明细清单.md》。" %
  (len(DET),
   sum(1 for d in DET if d["类别"].strip().startswith("A")),
   sum(1 for d in DET if d["类别"].strip().startswith("B")),
   sum(1 for d in DET if d["类别"].strip().startswith("C")),
   sum(1 for d in DET if d["类别"].strip().startswith("D")),
   len(Q), len(X), len(CONF)), 9)

doc.add_page_break()
H("一、一级问题（须在会审会上优先解决，共 %d 项）" % len(Q), 13)
rows = []
for q in Q:
    body = q["body"]
    m = re.search(r"\*\*建议处理\*\*：(.*)", body, re.S)
    sug = clean(m.group(1)) if m else ""
    desc = re.sub(r"\*\*建议处理\*\*.*", "", body, flags=re.S).strip()
    desc = desc.replace("**涉及图纸**：", "涉及图纸：")
    rows.append([q["id"], "%s·%s" % (q["cls"], q["sev"]), q["title"], desc, sug, ""])
mktable(["编号", "类别·程度", "问题标题", "问题描述与依据", "建议处理措施", "设计答复/处理意见"],
        [1.6, 1.8, 4.2, 9.4, 6.0, 4.0], rows, fsize=6.5)

doc.add_page_break()
H("二、跨专业矛盾专题（共 %d 项）" % len(X), 13)
rows = [[x["id"], x["title"], x["body"], ""] for x in X]
mktable(["编号", "议题", "问题描述与依据", "设计答复/处理意见"], [1.6, 4.0, 15.4, 6.0], rows, fsize=6.5)

doc.add_page_break()
H("三、需设计单位书面确认／补充事项（共 %d 项）" % len(CONF), 13)
rows = [[r[0], r[1], r[2], r[3], "", ""] for r in CONF]
mktable(["序号", "事项", "对应问题编号", "要求交付物", "设计答复", "交付日期"],
        [1.2, 6.0, 2.4, 7.4, 6.0, 4.0], rows, fsize=7.5)

doc.add_page_break()
H("四、各专业问题明细（共 %d 条）" % len(DET), 13)
rows = [[d["统一编号"], d["专业"], d["图号图名"], d["位置"], d["问题描述"], d["类别"], d["严重程度"], d["建议处理措施"], ""]
        for d in DET]
mktable(["编号", "专业", "图号/图名", "位置", "问题描述", "类别", "程度", "建议处理措施", "设计答复"],
        [1.5, 1.5, 3.4, 2.6, 7.6, 0.9, 0.9, 4.4, 4.2], rows, fsize=6)

DOCX = os.path.join(ROOT, "04_图纸会审记录表.docx")
doc.save(DOCX)
print("saved", DOCX)
